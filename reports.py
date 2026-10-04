"""
Stream-table export (CSV / Excel) for the DWSIM MCP tools.

Takes rows from stream_table (already converted to a unit system) and lays
them out the way process engineers expect: one column per stream, one row
per property, followed by each compound's mole and mass fraction.
"""

import csv
import os
from typing import Any, Dict, List


def _matrix(rows: List[Dict[str, Any]]):
    streams = [r["name"] for r in rows]
    scalar_keys = []
    for r in rows:
        for k, v in r.items():
            if k in ("name", "mole_fractions", "mass_fractions", "phases") or isinstance(v, dict):
                continue
            if k not in scalar_keys:
                scalar_keys.append(k)
    compounds = []
    for r in rows:
        for c in (r.get("mole_fractions") or {}):
            if c not in compounds:
                compounds.append(c)

    table = [["Property"] + streams]
    for k in scalar_keys:
        table.append([k] + [r.get(k) for r in rows])
    for basis, key in (("mole fraction", "mole_fractions"), ("mass fraction", "mass_fractions")):
        for c in compounds:
            table.append([f"{basis}: {c}"] + [(r.get(key) or {}).get(c) for r in rows])
    return table


def write_stream_table(rows: List[Dict[str, Any]], file_path: str) -> Dict[str, Any]:
    if not rows:
        raise ValueError("There are no material streams to export.")
    return _write(_matrix(rows), file_path, "Streams", transpose_header=True)


def write_table(rows: List[Dict[str, Any]], file_path: str, sheet: str = "Results") -> Dict[str, Any]:
    """Write a list of flat dicts as a normal table (one row per dict)."""
    if not rows:
        raise ValueError("Nothing to export.")
    cols = []
    for r in rows:
        for k in r:
            if k not in cols:
                cols.append(k)
    table = [cols] + [[r.get(c) for c in cols] for r in rows]
    return _write(table, file_path, sheet, transpose_header=False)


def _write(table, file_path: str, sheet: str, transpose_header: bool) -> Dict[str, Any]:
    os.makedirs(os.path.dirname(os.path.abspath(file_path)) or ".", exist_ok=True)
    ext = os.path.splitext(file_path)[1].lower()
    if ext == ".csv":
        with open(file_path, "w", newline="", encoding="utf-8-sig") as f:
            csv.writer(f).writerows(table)
    elif ext == ".xlsx":
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        wb = Workbook()
        ws = wb.active
        ws.title = sheet
        for row in table:
            ws.append(row)
        header_fill = PatternFill("solid", fgColor="DDEBF7")
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", wrap_text=not transpose_header)
        if transpose_header:
            for row in ws.iter_rows(min_row=2, max_col=1):
                row[0].font = Font(bold=True)
        for row in ws.iter_rows(min_row=2, min_col=2 if transpose_header else 1):
            for cell in row:
                if isinstance(cell.value, float):
                    cell.number_format = "0.0000" if abs(cell.value) < 10 else "#,##0.00"
        for i in range(1, len(table[0]) + 1):
            width = max(len(str(r[i - 1])) for r in table if r[i - 1] is not None) if i == 1 and transpose_header \
                else max(14, min(40, len(str(table[0][i - 1])) + 2))
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = width + 2
        ws.freeze_panes = "B2" if transpose_header else "A2"
        wb.save(file_path)
    else:
        raise ValueError("file_path must end in .csv or .xlsx.")
    return {"file_path": os.path.abspath(file_path), "columns": len(table[0]) - (1 if transpose_header else 0),
            "rows": len(table) - 1}
